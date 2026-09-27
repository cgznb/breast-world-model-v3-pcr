"""Reconcile original and revised pCR results without fitting or inference."""

from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ.update(OMP_NUM_THREADS="1", OPENBLAS_NUM_THREADS="1", MKL_NUM_THREADS="1")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.stats import rankdata
from sklearn.metrics import roc_auc_score

from scripts.run_full978_anti_overfit import _atomic_csv, _atomic_text
from src.first_post_pcr_data import identity, now, read_json, write_json

ROOT = Path(__file__).resolve().parents[1]
V1 = ROOT / "results/registered_three_phase_roi32_sequential_pcr_v1_20260918"
PCR = ROOT / "results/registered_three_phase_roi32_pcr_v1_20260917"
V3 = ROOT / "results/registered_three_phase_roi32_fixed_pcr_v3_20260918"
V4 = ROOT / "results/registered_three_phase_roi32_overfit_ablation_v4_20260918"
LATEST = V4 / "candidate_followup"
OUT = ROOT / "results/registered_three_phase_pcr_version_comparison_20260918"
WINDOWS = ["T0", "T0-T1", "T0-T2", "T0-T3"]
SEEDS = list(range(42, 52))
SOURCES = ["real", "direct_mc4", "rollout_mc4", "previous_real_mc4"]
LABELS = {"real": "\u771f\u5b9e\u5f71\u50cf", "direct_mc4": "T0\u76f4\u63a5\u751f\u6210", "rollout_mc4": "Rollout",
          "previous_real_mc4": "\u771f\u5b9e\u524d\u4e00\u9636\u6bb5\u751f\u6210"}
VERSIONS = {
    "v1_original": {"title": "\u7b2c\u4e00\u7248 V1", "windows": WINDOWS,
                    "recipe": "\u539f\u59cb\u914d\u7f6e\uff0c\u6309\u5404\u81ea\u9a8c\u8bc1\u6298\u9009\u62e9\u6700\u4f73\u68c0\u67e5\u70b9", "prefix": "V1"},
    "v3_regularized": {"title": "\u8f83\u65e9\u7684\u5168\u7a97\u53e3\u6539\u826f V3", "windows": WINDOWS,
                        "recipe": "\u7edf\u4e0040\u8f6e\uff1b\u7ec4\u5408\u6b63\u5219\u5316\u65b9\u6848\uff0c\u5c5e\u4e8e\u8f83\u65e9\u7684\u6539\u826f\u5b9e\u9a8c", "prefix": "V3"},
    "v4_latest": {"title": "\u6700\u65b0\u5c40\u90e8\u6539\u826f V4", "windows": ["T0-T3"],
                  "recipe": "\u4ec5T0-T3\uff0c40\u8f6e\uff0c\u4fdd\u7559\u5e73\u8861BCE\u5e76\u52a0\u6b8b\u5deeL2\u60e9\u7f5a0.1", "prefix": "V4"},
    "fixed40_reference": {"title": "\u56fa\u5b9a40\u8f6e\u53c2\u8003\u7ec4", "windows": WINDOWS,
                          "recipe": "\u7528\u4e8e\u63a7\u5236\u8bad\u7ec3\u9884\u7b97\u7684\u53c2\u8003\u7ec4\uff1b\u4e0e\u6700\u65e9V1\u4e0d\u662f\u540c\u4e00\u6279\u6743\u91cd", "prefix": "Ref40"},
}
KEYS = ["window", "seed", "source"]
FOLDS = [f"F{i}_auroc" for i in range(1, 6)]
INPUTS = {}


def csv(path, **kwargs):
    path = Path(path)
    INPUTS[str(path)] = identity(path)
    return pd.read_csv(path, float_precision="round_trip", **kwargs)


def document(path):
    path = Path(path)
    INPUTS[str(path)] = identity(path)
    return read_json(path)


def verify_inventory(entries):
    for path, expected in entries.items():
        if identity(path) != expected:
            raise ValueError(f"Source artifact changed: {path}")


def rank_auc(labels, probability):
    labels, probability = np.asarray(labels), np.asarray(probability, dtype=np.float64)
    if set(np.unique(labels)) != {0, 1} or not np.isfinite(probability).all():
        raise ValueError("Invalid labels or probabilities")
    if ((probability < 0) | (probability > 1)).any():
        raise ValueError("Probabilities outside [0,1]")
    positive = int(labels.sum())
    ranks = rankdata(probability, method="average")
    value = (ranks[labels == 1].sum() - positive * (positive + 1) / 2) / (positive * (len(labels) - positive))
    if abs(value - roc_auc_score(labels, probability)) > 1e-12:
        raise ValueError("Independent rank AUROC disagrees with sklearn")
    return float(value)


def audit(version, summary, probabilities, ensemble, expected_ids):
    expected = {(w, seed, source) for w in VERSIONS[version]["windows"] for seed in SEEDS for source in SOURCES}
    if len(summary) != len(expected) or set(summary[KEYS].itertuples(index=False, name=None)) != expected:
        raise ValueError("Missing/duplicate window, input or seed")
    probabilities = probabilities[probabilities.source.isin(SOURCES)].copy()
    ensemble = ensemble[ensemble.source.isin(SOURCES)].copy()
    if probabilities.duplicated(KEYS + ["fold", "patient_id"]).any() or ensemble.duplicated(KEYS + ["patient_id"]).any():
        raise ValueError("Duplicate model/patient probability")
    if set(probabilities.groupby(KEYS).groups) != expected or set(ensemble.groupby(KEYS).groups) != expected:
        raise ValueError("Prediction and summary coverage differ")
    lookup = summary.set_index(KEYS)
    raw_groups, ensemble_groups = probabilities.groupby(KEYS), ensemble.groupby(KEYS)
    rows, max_error, mean_error = [], 0.0, 0.0
    for key in sorted(expected):
        part, saved = raw_groups.get_group(key), lookup.loc[key]
        scores, labels_by_fold, probability_by_fold = [], [], []
        for fold in range(1, 6):
            batch = part[part.fold == fold].sort_values("patient_id")
            if batch.patient_id.tolist() != expected_ids or len(batch) != 102 or int(batch.label.sum()) != 32:
                raise ValueError("Historical patient population or outcome coverage changed")
            labels_by_fold.append(batch.label.to_numpy())
            probability_by_fold.append(batch.probability.to_numpy())
            scores.append(rank_auc(batch.label, batch.probability))
        if len(part) != 510 or not all(np.array_equal(labels_by_fold[0], y) for y in labels_by_fold):
            raise ValueError("Five models must evaluate the same outcomes")
        ep = ensemble_groups.get_group(key).sort_values("patient_id")
        if ep.patient_id.tolist() != expected_ids or not np.array_equal(ep.label.to_numpy(), labels_by_fold[0]):
            raise ValueError("Ensemble patients or labels changed")
        ensemble_auc = rank_auc(ep.label, ep.probability)
        probability_error = float(np.max(np.abs(np.mean(probability_by_fold, axis=0, dtype=np.float64) - ep.probability.to_numpy())))
        mean_error = max(mean_error, probability_error)
        if probability_error > (1e-7 if version == "v1_original" else 1e-12):
            raise ValueError("Saved ensemble differs from mean of patient probabilities")
        actual = np.array([*scores, np.mean(scores), np.std(scores, ddof=1), ensemble_auc])
        recorded = saved[[*FOLDS, "auroc_mean", "auroc_fold_sd", "probability_ensemble_auroc"]].to_numpy(dtype=float)
        max_error = max(max_error, float(np.max(np.abs(actual - recorded))))
        if max_error > 1e-12:
            raise ValueError("Saved summary disagrees with individual patient predictions")
        rows.append({"version": version, "window": key[0], "seed": key[1], "source": key[2],
                     "patients": 102, "positive": 32, "folds": 5,
                     **dict(zip(FOLDS, scores)), "auroc_mean": float(np.mean(scores)),
                     "auroc_fold_sd": float(np.std(scores, ddof=1)), "probability_ensemble_auroc": ensemble_auc})
    result = pd.DataFrame(rows)
    for window, sources in (("T0", SOURCES), ("T0-T1", SOURCES[1:])):
        subset = result[(result.window == window) & result.source.isin(sources)]
        if len(subset):
            for _, values in subset.groupby("seed"):
                matrix = values[[*FOLDS, "probability_ensemble_auroc"]].to_numpy()
                if np.max(np.abs(matrix - matrix[0])) > 1e-12:
                    raise ValueError("Shared T0 or first-transition results differ")
    return result, {"summaries": len(result), "fold_auc_replays": len(result) * 5,
                    "ensemble_auc_replays": len(result), "summary_max_error": max_error,
                    "ensemble_probability_max_error": mean_error}


def original(expected_ids):
    report = document(V1 / "evaluation/per_seed_input_folds/verification.json")
    assert report["passed"]
    verify_inventory(report["input_identities"])
    summary = csv(V1 / "evaluation/per_seed_input_folds/holdout_seed_summary.csv").rename(columns={"depth": "window"})
    summary = summary[summary.source.isin(SOURCES)]
    raw = csv(V1 / "evaluation/fold_predictions.csv", dtype={"patient_id": str, "probability": np.float32})
    raw = raw.rename(columns={"temporal_depth": "window"})
    keys = ["window", "seed", "fold", "patient_id", "label"]
    branches = [raw[raw.source == "real"][[*keys, "source", "probability"]]]
    for mode in ("direct", "rollout", "previous_real"):
        draws = raw[raw.source.str.startswith(mode + "_draw_")]
        if not (draws.groupby(keys).source.nunique() == 4).all():
            raise ValueError("MC4 requires four matched generation draws")
        part = draws.groupby(keys, as_index=False).probability.mean()
        part["source"] = mode + "_mc4"
        branches.append(part)
    probabilities = pd.concat(branches, ignore_index=True)
    probabilities["fold"] += 1
    ensemble = csv(V1 / "evaluation/predictions.csv", dtype={"patient_id": str, "probability": np.float32})
    ensemble = ensemble.rename(columns={"temporal_depth": "window"})
    return audit("v1_original", summary, probabilities, ensemble, expected_ids)


def revised(root, model_versions, expected_ids):
    assert document(root / "COMPLETE.json")["passed"]
    verify_inventory(document(root / "FROZEN_MODELS.json")["artifacts"])
    evaluation = root / "evaluation/holdout"
    summary = csv(evaluation / "seed_fivefold_summary.csv")
    summary = summary.rename(columns={f"fold_{i}_auroc": f"F{i}_auroc" for i in range(1, 6)})
    em = csv(evaluation / "probability_ensemble_metrics.csv")
    em = em.rename(columns={"auroc": "probability_ensemble_auroc"})
    keys = ["model", "depth", "seed", "source"]
    summary = summary.merge(em[[*keys, "probability_ensemble_auroc"]], on=keys, validate="one_to_one")
    probabilities = csv(evaluation / "mc4_fold_predictions.csv", dtype={"patient_id": str})
    probabilities["fold"] = probabilities.outer + 1
    probabilities["window"] = probabilities.depth.map(dict(enumerate(WINDOWS, 1)))
    ensemble = csv(evaluation / "probability_ensemble_predictions.csv", dtype={"patient_id": str})
    ensemble["window"] = ensemble.depth.map(dict(enumerate(WINDOWS, 1)))
    result, audits = [], {}
    for model, version in model_versions.items():
        frame, evidence = audit(version, summary[summary.model == model], probabilities[probabilities.model == model],
                                ensemble[ensemble.model == model], expected_ids)
        result.append(frame)
        audits[version] = evidence
    return pd.concat(result, ignore_index=True), audits


def markdown_version(frame, version):
    info = VERSIONS[version]
    lines = ["# " + info["title"], "", info["recipe"] + "\u3002", "",
             "\u4ee5\u4e0b\u5747\u5728\u540c\u4e00\u6279102\u540d\u5386\u53f2\u7559\u51fa\u60a3\u8005\u4e0a\u8bc4\u5206\uff0832\u540d\u9633\u6027\uff09\uff1b\u5404\u79cd\u5b50\u5355\u72ec\u5c55\u793a\uff0c\u672a\u8de8\u79cd\u5b50\u5e73\u5747\u3002",
             "\u751f\u6210\u8f93\u5165\u5148\u5728\u6bcf\u4e2a\u6298\u6a21\u578b\u5185\u5e73\u57474\u6761\u5b8c\u6574\u751f\u6210\u5e8f\u5217\u7684\u9884\u6d4b\u6982\u7387\uff08MC4\uff09\u3002", "",
             "- \u4e94\u6298AUC\uff1a\u5148\u5206\u522b\u8ba1\u7b97\u4e94\u4e2a\u6298\u6a21\u578b\u7684AUC\uff0c\u518d\u62a5\u544a\u5747\u503c\u53ca\u6298\u95f4\u6837\u672c\u6807\u51c6\u5dee\u3002",
             "- \u6982\u7387\u96c6\u6210AUC\uff1a\u5148\u5e73\u5747\u4e94\u4e2a\u6298\u6a21\u578b\u5bf9\u6bcf\u4f4d\u60a3\u8005\u7684\u9884\u6d4b\u6982\u7387\uff0c\u518d\u8ba1\u7b97\u4e00\u4e2aAUC\u3002",
             "- 102\u4eba\u66fe\u7528\u4e8e\u751f\u6210\u6a21\u578b\u9a8c\u8bc1\uff0c\u56e0\u6b64\u8fd9\u662f\u63a2\u7d22\u6027\u5386\u53f2\u8bc4\u4f30\u3002\u6298\u95f4\u6807\u51c6\u5dee\u4e0d\u662f\u60a3\u8005\u7f6e\u4fe1\u533a\u95f4\u3002", ""]
    if version == "v4_latest":
        lines += ["\u672c\u8f6e\u53ea\u8bad\u7ec3\u4e86T0-T3\uff1bT0\u3001T0-T1\u3001T0-T2\u6ca1\u6709\u672c\u8f6e\u65b0\u6a21\u578b\uff0c\u4e0d\u586b\u5165\u65e7\u7248\u672c\u7684\u6570\u503c\u3002", ""]
    indexed = frame.set_index(KEYS)
    for window in info["windows"]:
        lines += ["## " + window, "", "### \u4e94\u6298AUC\u5747\u503c\u4e0e\u6807\u51c6\u5dee", "",
                  "| \u79cd\u5b50 | " + " | ".join(LABELS[s] for s in SOURCES) + " |", "|---|---:|---:|---:|---:|"]
        for seed in SEEDS:
            values = [indexed.loc[(window, seed, s)] for s in SOURCES]
            lines.append(f"| {seed} | " + " | ".join(f"{r.auroc_mean:.5f} \u00b1 {r.auroc_fold_sd:.5f}" for r in values) + " |")
        lines += ["", "### \u4e94\u6298\u9884\u6d4b\u6982\u7387\u5e73\u5747\u540e\u7684AUC", "",
                  "| \u79cd\u5b50 | " + " | ".join(LABELS[s] for s in SOURCES) + " |", "|---|---:|---:|---:|---:|"]
        for seed in SEEDS:
            lines.append(f"| {seed} | " + " | ".join(f"{indexed.loc[(window, seed, s)].probability_ensemble_auroc:.5f}" for s in SOURCES) + " |")
        lines.append("")
    lines += ["T0\u5404\u8f93\u5165\u76f8\u540c\uff0c\u56e0\u4e3a\u53ea\u4f7f\u7528\u771f\u5b9eT0\uff1bT0-T1\u4e09\u79cd\u751f\u6210\u8f93\u5165\u76f8\u540c\uff0c\u56e0\u4e3a\u5171\u4eab\u7b2c\u4e00\u6b21\u751f\u6210\u3002", "",
              "[\u6c47\u603bExcel](pcr_versions.xlsx) | [\u5168\u90e8F1-F5\u53ca\u4e24\u79cd\u7edf\u8ba1](all_versions.csv) | [\u7248\u672c\u8bf4\u660e](README.md)", ""]
    _atomic_text(OUT / f"{version}.md", "\n".join(lines))


def compact_tables(all_results):
    lines = ["\u6bcf\u683c\u4e3a\uff1a\u4e94\u6298AUC\u5747\u503c / \u4e94\u6298\u6982\u7387\u96c6\u6210AUC\u3002", ""]
    first = all_results[all_results.version == "v1_original"].set_index(KEYS)

    def cell(indexed, window, seed, source):
        row = indexed.loc[(window, seed, source)]
        return f"{row.auroc_mean:.5f} / {row.probability_ensemble_auroc:.5f}"

    lines += ["\u7b2c\u4e00\u7248\uff1aT0\u4e0eT0-T1", "", "| \u79cd\u5b50 | T0\uff1a\u6240\u6709\u8f93\u5165 | T0-T1\uff1a\u771f\u5b9e | T0-T1\uff1a\u4e09\u79cd\u751f\u6210 |", "|---|---:|---:|---:|"]
    for seed in SEEDS:
        lines.append(f"| {seed} | " + " | ".join([cell(first, "T0", seed, "real"), cell(first, "T0-T1", seed, "real"),
                                                      cell(first, "T0-T1", seed, "direct_mc4")]) + " |")
    for version, window in (("v1_original", "T0-T2"), ("v1_original", "T0-T3"), ("v4_latest", "T0-T3")):
        indexed = all_results[all_results.version == version].set_index(KEYS)
        lines += ["", VERSIONS[version]["title"] + "\uff1a" + window, "",
                  "| \u79cd\u5b50 | " + " | ".join(LABELS[s] for s in SOURCES) + " |", "|---|---:|---:|---:|---:|"]
        for seed in SEEDS:
            lines.append(f"| {seed} | " + " | ".join(cell(indexed, window, seed, s) for s in SOURCES) + " |")
    _atomic_text(OUT / "compact_tables.md", "\n".join(lines) + "\n")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    cohort = document(PCR / "cohort.json")
    ids = sorted(cohort["split"]["val"])
    if len(ids) != 102:
        raise ValueError("Historical evaluation population changed")
    original_results, original_audit = original(ids)
    fixed_results, fixed_audit = revised(V3, {"primary": "v3_regularized", "reference": "fixed40_reference"}, ids)
    latest_results, latest_audit = revised(LATEST, {"primary": "v4_latest"}, ids)
    all_results = pd.concat([original_results, fixed_results, latest_results], ignore_index=True)
    availability = pd.DataFrame([{"version": version, "title": meta["title"], "window": window,
                                  "available": window in meta["windows"], "recipe": meta["recipe"]}
                                 for version, meta in VERSIONS.items() for window in WINDOWS])
    deltas = []
    for left, right in (("v3_regularized", "v1_original"), ("v4_latest", "v1_original"),
                        ("v4_latest", "fixed40_reference")):
        a = all_results[all_results.version == left].set_index(KEYS)
        b = all_results[all_results.version == right].set_index(KEYS).loc[a.index]
        delta = a[["auroc_mean", "probability_ensemble_auroc"]] - b[["auroc_mean", "probability_ensemble_auroc"]]
        delta = delta.rename(columns={"auroc_mean": "delta_fold_mean_auc", "probability_ensemble_auroc": "delta_probability_ensemble_auc"})
        delta["left_version"], delta["right_version"] = left, right
        deltas.append(delta.reset_index())
    deltas = pd.concat(deltas, ignore_index=True)
    development = csv(V1 / "evaluation/per_seed_input_folds/development_seed_summary.csv").rename(columns={"depth": "window"})
    development["version"] = "v1_original"
    earlier_dev = csv(V3 / "evaluation/development/seed_fivefold_summary.csv")
    earlier_dev["version"] = earlier_dev.model.map({"primary": "v3_regularized", "reference": "fixed40_reference"})
    earlier_dev = earlier_dev.rename(columns={f"fold_{i}_auroc": f"F{i}_auroc" for i in range(1, 6)})
    newest_dev = csv(V4 / "per_seed_fivefold.csv")
    newest_dev = newest_dev[newest_dev.arm == "residual_penalty"].rename(
        columns={"validation_auroc_mean": "auroc_mean", "validation_auroc_fold_sd": "auroc_fold_sd"})
    newest_dev["version"] = "v4_latest"
    columns = ["version", "window", "seed", *FOLDS, "auroc_mean", "auroc_fold_sd"]
    development = pd.concat([development[columns], earlier_dev[columns], newest_dev[columns]], ignore_index=True)
    for frame in (all_results, development):
        values = frame[FOLDS].to_numpy(dtype=float)
        np.testing.assert_allclose(values.mean(axis=1), frame.auroc_mean, atol=1e-12, rtol=0)
        np.testing.assert_allclose(values.std(axis=1, ddof=1), frame.auroc_fold_sd, atol=1e-12, rtol=0)
    tables = {"\u7248\u672c\u4e0e\u8986\u76d6\u8303\u56f4": availability, "\u8de8\u7248\u672c\u5dee\u503c": deltas, "\u5f00\u53d1\u96c6\u771f\u5b9e\u5f71\u50cf": development}
    for version, meta in VERSIONS.items():
        frame = all_results[all_results.version == version].copy()
        markdown_version(frame, version)
        for key, suffix in (("auroc_mean", "\u4e94\u6298\u5747\u503c"), ("probability_ensemble_auroc", "\u6982\u7387\u96c6\u6210")):
            table = frame.pivot(index=["window", "seed"], columns="source", values=key)[SOURCES].reset_index()
            table.columns.name = None
            tables[meta["prefix"] + "_" + suffix] = table.rename(columns=LABELS)
        tables[meta["prefix"] + "_F1-F5"] = frame
    _atomic_csv(OUT / "all_versions.csv", all_results)
    _atomic_csv(OUT / "version_differences.csv", deltas)
    _atomic_csv(OUT / "availability.csv", availability)
    _atomic_csv(OUT / "development_real_only.csv", development)
    with pd.ExcelWriter(OUT / "pcr_versions.xlsx", engine="openpyxl") as writer:
        for name, frame in tables.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "C2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="294A42")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[1].height = 32
            for index, column in enumerate(frame.columns, 1):
                width = 25 if column in ("version", "source", "left_version", "right_version") else 19
                sheet.column_dimensions[get_column_letter(index)].width = 80 if column == "recipe" else 32 if column == "title" else width
                if pd.api.types.is_float_dtype(frame[column]):
                    for cells in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                        for cell in cells:
                            cell.number_format = "0.00000"
    for name, expected in tables.items():
        actual = pd.read_excel(OUT / "pcr_versions.xlsx", sheet_name=name)
        pd.testing.assert_frame_equal(actual, expected.reset_index(drop=True), check_dtype=False, atol=1e-12, rtol=1e-12)
    compact_tables(all_results)
    lines = ["# pCR\u7248\u672c\u5bf9\u6bd4\uff1a\u4e24\u79cd\u4e94\u6298\u7edf\u8ba1", "",
             "[\u5b8c\u6574Excel](pcr_versions.xlsx) | [\u7cbe\u7b80\u603b\u8868](compact_tables.md) | [\u6240\u6709\u9010\u6298\u6570\u636e](all_versions.csv)", "",
             "## \u7248\u672c\u5206\u5f00\u770b", "", "| \u7248\u672c | \u5b8c\u6210\u7684\u65f6\u95f4\u7a97 | \u8bad\u7ec3\u4e0e\u9009\u62e9 | \u8be6\u7ec6\u7ed3\u679c |", "|---|---|---|---|"]
    for version, meta in VERSIONS.items():
        lines.append(f"| {meta['title']} | {', '.join(meta['windows'])} | {meta['recipe']} | [{meta['prefix']}\u8868\u683c]({version}.md) |")
    lines += ["", "## \u7edf\u8ba1\u53e3\u5f84", "",
              "\u4e3b\u8868\u90fd\u662f\u540c\u4e00\u6279102\u540d\u5386\u53f2\u7559\u51fa\u60a3\u8005\uff0832\u540d\u9633\u6027\uff09\uff0c\u6bcf\u4e2a\u79cd\u5b50\u5355\u72ec\u7edf\u8ba1\uff0c\u4e94\u4e2a\u6a21\u578b\u5747\u5bf9\u8fd9102\u4eba\u9884\u6d4b\u3002",
              "\u4e94\u6298AUC\u5747\u503c=\u4e94\u4e2a\u6a21\u578b\u5206\u522b\u8ba1\u7b97AUC\u540e\u518d\u5e73\u5747\uff1b\u6298\u95f4SD\u53ea\u53cd\u6620\u6a21\u578b\u6ce2\u52a8\u3002",
              "\u6982\u7387\u96c6\u6210AUC=\u5148\u5e73\u5747\u4e94\u4e2a\u6a21\u578b\u5bf9\u540c\u4e00\u60a3\u8005\u7684\u9884\u6d4b\u6982\u7387\uff0c\u518d\u8ba1\u7b97\u4e00\u4e2aAUC\u3002AUC\u975e\u7ebf\u6027\uff0c\u4e24\u8005\u901a\u5e38\u4e0d\u540c\u3002",
              "\u751f\u6210\u8f93\u5165\u6cbf\u7528MC4\uff1a\u6bcf\u4e2a\u6a21\u578b\u5185\u5e73\u57474\u6761\u5b8c\u6574\u751f\u6210\u5e8f\u5217\u7684pCR\u6982\u7387\u3002\u539f\u59cbV1\u4fdd\u7559\u539f\u4fdd\u5b58\u7cbe\u5ea6\u548c\u805a\u5408\u7ed3\u679c\u3002",
              "\u6240\u6709\u8f93\u5165\u4fdd\u7559\u771f\u5b9eT0\uff1b\u76f4\u63a5\u751f\u6210\u4eceT0\u5206\u522b\u9884\u6d4b\u540e\u7eed\uff0crollout\u9010\u9636\u6bb5\u751f\u6210\uff0c\u771f\u5b9e\u524d\u4e00\u9636\u6bb5\u751f\u6210\u4f7f\u7528\u771f\u5b9e\u524d\u6b21\u5f71\u50cf\u3002",
              "T0\u7684\u56db\u79cd\u8f93\u5165\u76f8\u540c\uff1bT0-T1\u7684\u4e09\u79cd\u751f\u6210\u8f93\u5165\u76f8\u540c\u3002", "",
              "## \u6bd4\u8f83\u8fb9\u754c", "",
              "\u6700\u65b0V4\u4ec5\u6709T0-T3\uff0c\u5176\u4f59\u4e09\u4e2a\u65f6\u95f4\u7a97\u6ca1\u6709\u672c\u8f6e\u65b0\u6a21\u578b\u3002V3\u662f\u8f83\u65e9\u7684\u5168\u7a97\u53e3\u5b9e\u9a8c\uff0c\u4e0d\u80fd\u628a\u5b83\u7684\u77ed\u7a97\u53e3\u62fc\u6210V4\u3002",
              "V4\u4e0eV1\u7684\u5dee\u5f02\u540c\u65f6\u5305\u542b\u8bad\u7ec3\u9884\u7b97\u548c\u68c0\u67e5\u70b9\u9009\u62e9\uff1b\u9694\u79bb\u6b8b\u5dee\u60e9\u7f5a\u7684\u5f71\u54cd\u5e94\u5bf9\u7167\u56fa\u5b9a40\u8f6e\u53c2\u8003\u7ec4\u3002",
              "\u8fd9\u4e9b\u6539\u826f\u7248\u662f\u5b9e\u9a8c\u5019\u9009\uff0c\u4e0d\u4ee3\u8868\u5df2\u7ecf\u4f18\u4e8e\u7b2c\u4e00\u7248\u3002\u8de8\u7248\u672c\u5dee\u503c\u53ea\u662f\u63cf\u8ff0\uff0c\u4e0d\u81ea\u52a8\u9009\u62e9\u7248\u672c\u6216\u79cd\u5b50\u3002",
              "102\u4eba\u66fe\u7528\u4e8e\u751f\u6210\u6a21\u578b\u9a8c\u8bc1\uff0c\u5c5e\u4e8e\u63a2\u7d22\u6027\u5386\u53f2\u8bc4\u4f30\uff1b\u4e0d\u80fd\u636e\u6b64\u58f0\u79f0\u72ec\u7acb\u7684\u5b8c\u6574\u6d41\u7a0b\u6cdb\u5316\u7ed3\u679c\u3002", "",
              "## \u5f00\u53d1\u96c6\u5355\u72ec\u4fdd\u5b58", "",
              "Excel\u7684\u5f00\u53d1\u96c6\u9875\u4ec5\u5217764\u4eba\u7684\u771f\u5b9e\u5f71\u50cf\u4e94\u6298\u9a8c\u8bc1\u3002\u6bcf\u4eba\u53ea\u6709\u4e00\u4e2a\u6298\u5916\u9884\u6d4b\uff0c\u4e0d\u62a5\u544a\u4f2a\u9020\u7684\u4e94\u6a21\u578b\u6982\u7387\u5e73\u5747\u9a8c\u8bc1AUC\u3002",
              "V1\u7684\u9a8c\u8bc1\u6298\u53c2\u4e0e\u68c0\u67e5\u70b9\u9009\u62e9\uff1bV3\u4e0e\u672c\u8f6eV4\u4f7f\u7528\u56fa\u5b9a40\u8f6e\u3002\u6ca1\u6709\u76f8\u5e94\u751f\u6210\u8f93\u5165\u7684\u5f00\u53d1\u96c6\u4e94\u6298\u7ed3\u679c\u3002", "",
              "[\u5f00\u53d1\u96c6\u771f\u5b9e\u5f71\u50cfCSV](development_real_only.csv) | [\u8de8\u7248\u672c\u5dee\u503c](version_differences.csv) | [\u6838\u9a8c\u8bb0\u5f55](verification.json)", ""]
    _atomic_text(OUT / "README.md", "\n".join(lines))
    verify_inventory(INPUTS)
    verification = {"passed": True, "checked_utc": now(), "versions": {"v1_original": original_audit, **fixed_audit, **latest_audit},
                    "patient_count": 102, "positive_count": 32, "classifier_seeds": SEEDS,
                    "historical_summaries": len(all_results), "development_summaries": len(development),
                    "workbook_sheets_roundtripped": len(tables), "across_seed_averaging": False,
                    "training_or_model_inference": False, "inputs_unchanged": True, "input_identities": INPUTS}
    write_json(OUT / "verification.json", verification)
    print({key: value for key, value in verification.items() if key != "input_identities"})
    print(OUT / "README.md")


if __name__ == "__main__":
    main()
