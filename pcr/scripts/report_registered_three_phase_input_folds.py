"""Report five individual model scores per seed, window, and imaging input."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from sklearn.metrics import average_precision_score, brier_score_loss, log_loss, roc_auc_score

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from scripts.plot_registered_three_phase_sequential_pcr import ARMS, DEPTHS, OUTPUT
from scripts.run_full978_anti_overfit import _atomic_csv, _atomic_text
from src.first_post_pcr_data import identity, now, read_json, write_json


ROOT = Path(__file__).resolve().parents[1]
PCR = ROOT / "results/registered_three_phase_roi32_pcr_v1_20260917"
DESTINATION = OUTPUT / "evaluation/per_seed_input_folds"
SEEDS = list(range(42, 52))
METRICS = ("auroc", "prauc", "logloss", "brier")
MAIN_ARMS = ("real", "direct_mc4", "rollout_mc4", "previous_real_mc4")


def scores(frame):
    y, p = frame.label.to_numpy(), frame.probability.to_numpy()
    assert np.isfinite(p).all() and ((p >= 0) & (p <= 1)).all() and set(y) == {0, 1}
    auc = float(roc_auc_score(y, p))
    positive, negative = p[y == 1, None], p[y == 0][None, :]
    rank_auc = float(((positive > negative) + .5 * (positive == negative)).mean())
    assert abs(auc - rank_auc) < 1e-12
    return {"patients": len(frame), "positives": int(y.sum()), "auroc": auc,
            "prauc": float(average_precision_score(y, p)),
            "logloss": float(log_loss(y, p, labels=[0, 1])),
            "brier": float(brier_score_loss(y, p))}, abs(auc - rank_auc)


def summarize(details):
    rows = []
    for (population, depth, seed, source), part in details.groupby(["population", "depth", "seed", "source"]):
        assert len(part) == 5 and set(part.fold) == set(range(1, 6))
        part = part.set_index("fold").sort_index()
        row = {"population": population, "depth": depth, "seed": int(seed), "source": source, "folds": 5}
        row.update({f"F{fold}_auroc": float(part.loc[fold, "auroc"]) for fold in range(1, 6)})
        for metric in METRICS:
            row[metric + "_mean"] = float(part[metric].mean())
            row[metric + "_fold_sd"] = float(part[metric].std(ddof=1))
        rows.append(row)
    return pd.DataFrame(rows)


def report():
    assert read_json(OUTPUT / "COMPLETE.json")["passed"]
    assert read_json(OUTPUT / "evaluation/independent_audit.json")["passed"]
    protocol = read_json(OUTPUT / "protocol.json")
    refs = protocol["classifiers"]
    expected_models = {(depth, seed, fold) for depth in DEPTHS for seed in SEEDS for fold in range(5)}
    assert len(refs) == 200 and {(r["depth"], r["seed"], r["fold"]) for r in refs} == expected_models
    inputs = {str(path): identity(path) for path in (OUTPUT / "protocol.json", OUTPUT / "COMPLETE.json",
              OUTPUT / "evaluation/independent_audit.json", OUTPUT / "evaluation/fold_predictions.csv",
              OUTPUT / "evaluation/predictions.csv", PCR / "cohort.json", PCR / "folds.json")}
    for ref in refs:
        assert identity(ref["path"]) == ref["identity"]
        Path(ref["path"]).relative_to(PCR)
        inputs[ref["path"]] = ref["identity"]
    cohort, folds = read_json(PCR / "cohort.json"), read_json(PCR / "folds.json")
    folds = {f["fold"]: f for f in folds}
    heldout, development = set(cohort["split"]["val"]), set(cohort["split"]["train"])
    assert len(heldout) == 102 and len(development) == 764 and not heldout & development
    raw = pd.read_csv(OUTPUT / "evaluation/fold_predictions.csv", dtype={"probability": np.float32})
    raw = raw.rename(columns={"temporal_depth": "depth"})
    keys = ["depth", "seed", "fold", "patient_id", "label"]
    assert len(raw) == 285600 and not raw.duplicated(["source", *keys]).any()
    assert (raw.groupby("patient_id").label.nunique() == 1).all()
    expected_sources = {"real", "copy_T0"} | {f"{mode}_draw_{draw}" for mode in
                        ("direct", "rollout", "previous_real") for draw in range(4)}
    assert set(raw.source) == expected_sources and set(raw.patient_id) == heldout
    for _, group in raw.groupby(["source", "depth", "seed", "fold"]):
        assert len(group) == 102 and set(group.patient_id) == heldout and group.label.sum() == 32
    branches = [raw[raw.source.isin(("real", "copy_T0"))][["source", *keys, "probability"]]]
    mc_error = 0.0
    for mode in ("direct", "rollout", "previous_real"):
        draws = raw[raw.source.str.startswith(mode + "_draw_")]
        assert (draws.groupby(keys).source.nunique() == 4).all()
        mean = draws.groupby(keys, as_index=False).probability.mean()
        independent = draws.pivot(index=keys, columns="source", values="probability").to_numpy(dtype=np.float64).mean(axis=1)
        mc_error = max(mc_error, float(np.abs(mean.set_index(keys).sort_index().probability.to_numpy() - independent).max()))
        mean["source"] = mode + "_mc4"
        branches.append(mean)
    assert mc_error < 1e-7
    probabilities = pd.concat(branches, ignore_index=True)
    assert len(probabilities) == 102000
    t0 = probabilities[probabilities.depth == "T0"].pivot(index=["seed", "fold", "patient_id"], columns="source", values="probability")
    assert float(t0.sub(t0.real, axis=0).abs().max().max()) == 0
    t1 = probabilities[probabilities.depth == "T0-T1"].pivot(index=["seed", "fold", "patient_id"], columns="source", values="probability")
    assert float(t1[list(MAIN_ARMS[1:])].sub(t1.direct_mc4, axis=0).abs().max().max()) == 0
    rows, max_auc_error = [], 0.0
    for (depth, seed, fold, source), frame in probabilities.groupby(["depth", "seed", "fold", "source"]):
        values, error = scores(frame)
        max_auc_error = max(max_auc_error, error)
        rows.append({"population": "historical_holdout_102", "depth": depth, "seed": int(seed),
                     "fold": int(fold) + 1, "source": source, **values})
    historical = pd.DataFrame(rows)
    assert len(historical) == 1000
    rows, validation_sets = [], {}
    for ref in refs:
        path = Path(ref["path"]).parent / "val_predictions.csv"
        inputs[str(path)] = identity(path)
        frame = pd.read_csv(path, dtype={"probability": np.float32})
        wanted = set(folds[ref["fold"]]["val_ids"])
        assert len(frame) == len(wanted) and frame.patient_id.is_unique and set(frame.patient_id) == wanted
        assert wanted <= development and not wanted & heldout
        assert set(frame.seed) == {ref["seed"]} and set(frame.fold) == {ref["fold"]}
        assert set(frame.temporal_depth) == {ref["depth"]}
        validation_sets.setdefault((ref["depth"], ref["seed"]), []).append(wanted)
        values, error = scores(frame)
        max_auc_error = max(max_auc_error, error)
        rows.append({"population": "development_validation_764", "depth": ref["depth"], "seed": ref["seed"],
                     "fold": ref["fold"] + 1, "source": "real", **values})
    for sets in validation_sets.values():
        assert len(sets) == 5 and set.union(*sets) == development and sum(map(len, sets)) == 764
    validation = pd.DataFrame(rows)
    assert len(validation) == 200
    historical_summary, validation_summary = summarize(historical), summarize(validation)
    assert len(historical_summary) == 200 and len(validation_summary) == 40
    old = pd.read_csv(OUTPUT / "evaluation/predictions.csv", dtype={"probability": np.float32})
    old = old[old.source.isin(ARMS)].rename(columns={"temporal_depth": "depth"})
    reference = []
    for (depth, seed, source), frame in old.groupby(["depth", "seed", "source"]):
        values, _ = scores(frame)
        reference.append({"depth": depth, "seed": seed, "source": source, "probability_ensemble_auroc": values["auroc"]})
    reference = pd.DataFrame(reference)
    historical_summary = historical_summary.merge(reference, on=["depth", "seed", "source"], validate="one_to_one")
    DESTINATION.mkdir(parents=True, exist_ok=True)
    tables = {"holdout_fold_metrics": historical, "holdout_seed_summary": historical_summary,
              "development_fold_metrics": validation, "development_seed_summary": validation_summary}
    export_error = 0.0
    for name, frame in tables.items():
        _atomic_csv(DESTINATION / (name + ".csv"), frame)
        if "summary" in name:
            saved = pd.read_csv(DESTINATION / (name + ".csv"))
            values = saved[[f"F{f}_auroc" for f in range(1, 6)]].to_numpy()
            export_error = max(export_error, float(np.abs(values.mean(1) - saved.auroc_mean).max()),
                               float(np.abs(values.std(1, ddof=1) - saved.auroc_fold_sd).max()))
    assert export_error < 1e-12
    write_reports(historical_summary, validation_summary)
    write_workbook(tables)
    plot(historical_summary)
    for path, expected in inputs.items():
        assert identity(path) == expected
    verification = {"passed": True, "checked_utc": now(), "source_run": str(PCR),
                    "evaluation_run": str(OUTPUT), "classifier_seeds": SEEDS, "folds": [1, 2, 3, 4, 5],
                    "holdout_fold_metric_rows": len(historical), "holdout_seed_window_input_summaries": len(historical_summary),
                    "development_fold_metric_rows": len(validation), "development_seed_window_summaries": len(validation_summary),
                    "raw_holdout_prediction_rows": len(raw), "MC4_probability_max_rounding_error": mc_error,
                    "independent_rank_auc_max_error": max_auc_error, "exported_mean_and_sd_max_error": export_error,
                    "all_T0_and_shared_T1_probability_error": 0.0, "inputs_and_models_unchanged": True,
                    "aggregation": "mean_and_sample_sd_of_five_model_AUROCs_per_seed; generated_MC4_probabilities_within_each_model_first",
                    "across_seed_averaging": False, "training_or_inference_performed": False,
                    "input_identities": inputs}
    write_json(DESTINATION / "verification.json", verification)
    print(json.dumps({k: v for k, v in verification.items() if k != "input_identities"}, indent=2))
    print(DESTINATION / "README.md")


def write_reports(historical, validation):
    lines = ["# Original ROI32 pCR v1: Five-Fold Results By Seed And Input", "",
             "Each cell is the mean AUROC +/- sample SD across five fold-trained models for ONE classifier seed. All ten seeds remain separate.",
             "For generated inputs, first average four generation-draw probabilities within each model, then calculate that model's AUROC.",
             "The primary four-input comparison evaluates each model on the SAME 102 historical holdout patients (32 positive). It is not five different validation subsets.",
             "These statistics differ from the earlier five-model probability-ensemble AUROC. That older statistic is retained separately in holdout_seed_summary.csv.",
             "The historical holdout participated in generator validation; fold SD is descriptive model variation, not a patient confidence interval.",
             "Real-predecessor generation uses observed follow-up images. Original ROI32 pCR v1 classifiers are used throughout; no new fitting or inference.",
             "Human fold numbers F1-F5 correspond to original IDs 0-4.", "",
             "[Excel workbook](fivefold_results.xlsx) | [F1-F5 details](fold_details.md) | [Holdout CSV](holdout_seed_summary.csv) | [Plot](fivefold_comparison.png)", ""]
    indexed = historical.set_index(["depth", "seed", "source"])
    for depth in DEPTHS:
        lines += ["## " + depth, "", "| Seed | Real | Direct T0 | Rollout | Real previous | Copy T0 |",
                  "|---:|---:|---:|---:|---:|---:|"]
        for seed in SEEDS:
            cells = []
            for arm in (*MAIN_ARMS, "copy_T0"):
                row = indexed.loc[(depth, seed, arm)]
                cells.append(f"{row.auroc_mean:.5f} +/- {row.auroc_fold_sd:.5f}")
            lines.append("| " + str(seed) + " | " + " | ".join(cells) + " |")
        lines.append("")
    lines += ["## Development Real-Input Validation", "",
              "These are the five disjoint validation subsets of the 764-patient development cohort, scored using real images only.",
              "The original v1 checkpoints were selected on these same fold-validation AUROCs, so this is selection validation, not an untouched outer test.",
              "Generated inputs have only been evaluated for the historical holdout in this original study. No generated development-CV score is inferred or fabricated.", "",
              "| Seed | T0 | T0-T1 | T0-T2 | T0-T3 |", "|---:|---:|---:|---:|---:|"]
    indexed = validation.set_index(["depth", "seed"])
    for seed in SEEDS:
        cells = [f"{indexed.loc[(depth, seed), 'auroc_mean']:.5f} +/- {indexed.loc[(depth, seed), 'auroc_fold_sd']:.5f}" for depth in DEPTHS]
        lines.append("| " + str(seed) + " | " + " | ".join(cells) + " |")
    _atomic_text(DESTINATION / "README.md", "\n".join(lines) + "\n")
    lines = ["# Individual F1-F5 AUROCs", "", "F1-F5 use human fold numbers 1-5. Mean and sample SD are calculated across these five AUCs, with no across-seed averaging.", ""]
    for title, summary in (("Same 102-patient historical holdout", historical), ("Real-only development validation", validation)):
        lines += ["## " + title, ""]
        for depth in DEPTHS:
            for arm, (label, _) in ARMS.items():
                part = summary[(summary.depth == depth) & (summary.source == arm)].sort_values("seed")
                if part.empty:
                    continue
                lines += [f"### {depth}: {label}", "", "| Seed | F1 | F2 | F3 | F4 | F5 | Mean | Fold SD |",
                          "|---:|---:|---:|---:|---:|---:|---:|---:|"]
                for row in part.itertuples():
                    values = [getattr(row, f"F{fold}_auroc") for fold in range(1, 6)] + [row.auroc_mean, row.auroc_fold_sd]
                    lines.append("| " + str(row.seed) + " | " + " | ".join(f"{value:.5f}" for value in values) + " |")
                lines.append("")
    _atomic_text(DESTINATION / "fold_details.md", "\n".join(lines) + "\n")


def write_workbook(tables):
    notes = pd.DataFrame([
        ("Study", "Original registered_three_phase_roi32_pcr_v1_20260917"),
        ("Primary population", "Same historical 102 patients for every model and input; 32 positives"),
        ("Summary", "Mean of five model AUROCs within each seed; SD is sample fold SD"),
        ("Generated inputs", "MC4 probabilities averaged within each fold model before its AUROC"),
        ("Previous table", "Five-model probability-ensemble AUROC is a separate reference column"),
        ("Fold numbering", "F1-F5 correspond to original fold IDs 0-4"),
        ("Seeds", "42-51 are classifier training seeds; no across-seed averaging"),
        ("Development tables", "Real-only disjoint fold validation across 764 patients; same validation selected checkpoints"),
        ("Limit", "Historical holdout was used for generator validation; real-previous uses real follow-up"),
    ], columns=["item", "description"])
    sheets = {"Notes": notes, "Holdout_F1_F5": tables["holdout_seed_summary"],
              "Holdout_Metrics": tables["holdout_fold_metrics"],
              "Development_F1_F5": tables["development_seed_summary"],
              "Development_Metrics": tables["development_fold_metrics"]}
    with pd.ExcelWriter(DESTINATION / "fivefold_results.xlsx", engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "E2" if name != "Notes" else "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="355C54")
                field = str(cell.value)
                sheet.column_dimensions[cell.column_letter].width = min(40, max(13, len(field) + 2))
            if name == "Notes":
                sheet.column_dimensions["B"].width = 110
            else:
                for index, field in enumerate(frame.columns, 1):
                    if frame[field].dtype.kind == "f":
                        for row in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                            row[0].number_format = "0.00000"
    for name, frame in sheets.items():
        reloaded = pd.read_excel(DESTINATION / "fivefold_results.xlsx", sheet_name=name)
        assert list(reloaded.columns) == list(frame.columns) and len(reloaded) == len(frame)
        for field in frame.columns:
            if frame[field].dtype.kind in "fiu":
                np.testing.assert_allclose(reloaded[field], frame[field], rtol=0, atol=1e-12)
            else:
                assert reloaded[field].tolist() == frame[field].tolist()


def plot(summary):
    plt.rcParams.update({"font.size": 10, "pdf.fonttype": 42, "savefig.facecolor": "white"})
    fig, axes = plt.subplots(2, 2, figsize=(12, 10))
    values = summary[summary.source.isin(MAIN_ARMS)].auroc_mean
    lower, upper = float(values.min()) - .005, float(values.max()) + .005
    for axis, depth in zip(axes.flat, DEPTHS):
        frame = summary[summary.depth == depth].pivot(index="seed", columns="source", values="auroc_mean")
        matrix = frame.loc[SEEDS, list(MAIN_ARMS)].to_numpy()
        graphic = axis.imshow(matrix, cmap="cividis", vmin=lower, vmax=upper, aspect="auto")
        axis.set(title=depth, xticks=range(4), xticklabels=["Real", "Direct T0", "Rollout", "Real previous"],
                 yticks=range(10), yticklabels=SEEDS, ylabel="Classifier seed")
        for row in range(10):
            for column in range(4):
                color = "white" if (matrix[row, column] - lower) / (upper - lower) < .48 else "black"
                axis.text(column, row, f"{matrix[row, column]:.4f}", ha="center", va="center", color=color, fontsize=9)
    fig.suptitle("Mean of five model AUROCs, separately for each classifier seed", fontsize=14, y=.97)
    fig.subplots_adjust(left=.07, right=.87, top=.9, bottom=.12, wspace=.26, hspace=.27)
    color_axis = fig.add_axes([.9, .16, .02, .68])
    fig.colorbar(graphic, cax=color_axis, label="Mean five-model AUROC")
    fig.text(.5, .065, "Same 102 historical holdout patients. Generated MC4 probabilities are averaged within each model before scoring.", ha="center", fontsize=9)
    fig.text(.5, .04, "F1-F5 and fold SD are in the tables. No across-seed averaging. This is distinct from averaging five model probabilities first.", ha="center", fontsize=9)
    for suffix in ("png", "pdf"):
        fig.savefig(DESTINATION / ("fivefold_comparison." + suffix), dpi=180, bbox_inches="tight")
    plt.close(fig)


if __name__ == "__main__":
    report()
